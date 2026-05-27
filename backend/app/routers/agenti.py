import httpx
import json
import mailparser
import pdfplumber
import openpyxl
import docx
import io
import re
from datetime import datetime, timedelta
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from sqlalchemy.orm import Session
from sqlalchemy import text
from pydantic import BaseModel
from typing import Optional
from anthropic import Anthropic
from app.database import get_mysql, get_sqlite, get_pratica
from app.config import AIS_BASE_URL
from app.auth import get_current_user
from sqlalchemy.ext.asyncio import AsyncSession
from app.email_service import check_new_emails_imap, mark_email_read_imap, send_email_smtp
from app.email_poller import get_inbox_status, get_inbox_emails, get_email_by_uid, mark_email_processed

router = APIRouter(prefix="/api/agenti", tags=["agenti"])

anthropic = Anthropic()


# ─── Helper: estrai testo da allegato ───────────────────────────────────────

def estrai_testo_allegato(nome: str, payload: bytes) -> str:
    ext = nome.lower().rsplit(".", 1)[-1] if "." in nome else ""
    try:
        if ext == "pdf":
            testo = []
            with pdfplumber.open(io.BytesIO(payload)) as pdf:
                for page in pdf.pages[:10]:
                    t = page.extract_text()
                    if t:
                        testo.append(t)
            return "\n".join(testo)[:4000]
        elif ext in ("docx", "doc"):
            d = docx.Document(io.BytesIO(payload))
            return "\n".join(p.text for p in d.paragraphs if p.text.strip())[:4000]
        elif ext in ("xlsx",):
            wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
            righe = []
            for ws in wb.worksheets[:3]:
                for row in ws.iter_rows(max_row=100, values_only=True):
                    riga = " | ".join(str(c) for c in row if c is not None)
                    if riga.strip():
                        righe.append(riga)
            return "\n".join(righe)[:4000]
    except Exception as e:
        return f"[Errore estrazione {nome}: {e}]"
    return ""


def ai(prompt: str, max_tokens: int = 800) -> str:
    msg = anthropic.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=max_tokens,
        messages=[{"role": "user", "content": prompt}]
    )
    return msg.content[0].text


def parse_json_safe(text: str) -> dict:
    text = text.strip()
    match = re.search(r'\{[\s\S]*\}', text)
    if match:
        raw = match.group(0)
    else:
        raw = text
    try:
        return json.loads(raw)
    except:
        try:
            open_b = raw.count('{') - raw.count('}')
            open_s = raw.count('[') - raw.count(']')
            fixed = raw + ']' * open_s + '}' * open_b
            return json.loads(fixed)
        except:
            return {"errore": "parsing fallito", "raw": raw[:300]}


def _lookup_gedoanag(mysql_db, nome: str = None, piva: str = None) -> dict | None:
    """Cerca un soggetto in gedoanag per nome o P.IVA. Restituisce None se non trovato."""
    try:
        if piva:
            row = mysql_db.execute(text(
                "SELECT id, ragi_socia, email, indirizzo, localita, provincia, part_iva "
                "FROM gedoanag WHERE part_iva = :piva AND (cancellato = 0 OR cancellato IS NULL) LIMIT 1"
            ), {"piva": piva}).fetchone()
            if row:
                return dict(row._mapping)
        if nome:
            row = mysql_db.execute(text(
                "SELECT id, ragi_socia, email, indirizzo, localita, provincia, part_iva "
                "FROM gedoanag WHERE ragi_socia LIKE :n AND (cancellato = 0 OR cancellato IS NULL) LIMIT 1"
            ), {"n": f"%{nome[:25]}%"}).fetchone()
            if row:
                return dict(row._mapping)
    except Exception:
        pass
    return None


async def _aggiorna_step(db: AsyncSession, pratica_id: int, step: int, extra: dict = None):
    """Aggiorna step_corrente (e opzionalmente altri campi) sulla pratica."""
    try:
        params = {"step": step, "id": pratica_id}
        set_clause = "step_corrente = :step"
        if extra:
            for k, v in extra.items():
                set_clause += f", {k} = :{k}"
                params[k] = v
        await db.execute(text(f"UPDATE pratiche SET {set_clause} WHERE id = :id"), params)
        await db.commit()
    except Exception:
        pass


async def _crea_approvazione(db: AsyncSession, pratica_id: int, step_numero: int, agente: str, output_ai: str):
    """Crea un record di approvazione pending per lo step AI completato."""
    try:
        await db.execute(text("""
            INSERT INTO approvazioni (pratica_id, step_numero, agente, output_ai, stato)
            VALUES (:pratica_id, :step_numero, :agente, :output_ai, 'pending')
        """), {
            "pratica_id": pratica_id,
            "step_numero": step_numero,
            "agente": agente,
            "output_ai": output_ai,
        })
        await db.commit()
    except Exception:
        pass


# ─── STEP 1a: Upload EML ────────────────────────────────────────────────────

@router.post("/step1/upload-eml")
async def step1_upload_eml(file: UploadFile = File(...), user=Depends(get_current_user)):
    contenuto = await file.read()
    try:
        mail = mailparser.parse_from_bytes(contenuto)
        corpo = ""
        if mail.text_plain:
            corpo = "\n".join(mail.text_plain)
        elif mail.text_html:
            corpo = re.sub(r'<[^>]+>', ' ', "\n".join(mail.text_html))
            corpo = re.sub(r'\s+', ' ', corpo).strip()
        else:
            corpo = mail.body or ""
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Errore parsing EML: {e}")

    allegati_estratti = []
    ESCLUDI_EXT = (".jpg", ".png", ".gif", ".jpeg", ".bmp", ".txt")
    for att in (mail.attachments or []):
        nome = att.get("filename", "") or att.get("mail_content_type", "allegato")
        if any(nome.lower().endswith(ext) for ext in ESCLUDI_EXT):
            allegati_estratti.append({"nome": nome, "testo": "(contenuto non estraibile)"})
            continue
        payload = att.get("payload")
        if payload and isinstance(payload, bytes):
            testo = estrai_testo_allegato(nome, payload)
        elif payload and isinstance(payload, str):
            try:
                import base64
                testo = estrai_testo_allegato(nome, base64.b64decode(payload))
            except:
                testo = ""
        else:
            testo = ""
        allegati_estratti.append({"nome": nome, "testo": testo if testo else "(contenuto non estraibile)"})

    return {
        "mittente": str(mail.from_),
        "destinatari": str(mail.to),
        "oggetto": mail.subject or "",
        "data": str(mail.date),
        "corpo": corpo[:6000],
        "allegati": allegati_estratti,
    }


# ─── INBOX EMAIL (nuovo sistema IMAP) ───────────────────────────────────────

@router.get("/inbox/status")
async def inbox_status(user=Depends(get_current_user)):
    """Restituisce lo stato dell'inbox email."""
    return get_inbox_status()


@router.get("/inbox/emails")
async def inbox_emails(user=Depends(get_current_user)):
    """Restituisce tutte le email ricevute."""
    return {"emails": get_inbox_emails()}


@router.get("/inbox/emails/{uid}")
async def inbox_email_detail(uid: str, user=Depends(get_current_user)):
    """Restituisce i dettagli di una email specifica."""
    email_data = get_email_by_uid(uid)
    if not email_data:
        raise HTTPException(status_code=404, detail="Email non trovata")
    return email_data


class ProcessEmailInput(BaseModel):
    uid: str

@router.post("/inbox/process")
async def inbox_process_email(body: ProcessEmailInput, user=Depends(get_current_user)):
    """Processa un'email dall'inbox per estrarre dati spedizione."""
    email_data = get_email_by_uid(body.uid)
    if not email_data:
        raise HTTPException(status_code=404, detail="Email non trovata")
    
    # Usa la stessa logica di step1/leggi-email ma con dati dall'inbox
    ESCLUDI = ("(contenuto non estraibile)", "[Errore estrazione")
    testo_allegati = ""
    for a in (email_data.get("allegati") or []):
        if not isinstance(a, dict):
            continue
        testo = a.get("testo", "")
        nome = a.get("nome", "")
        if any(testo.startswith(e) for e in ESCLUDI):
            continue
        testo_allegati += f"\n\n--- ALLEGATO: {nome} ---\n{testo[:2000]}"
    
    risposta = ai(f"""Sei un operatore di PF Ship Srl, agenzia spedizioni di Napoli.
Hai ricevuto questa email da un cliente italiano o dal tuo agente in Cina con un ordine di spedizione.
Devi estrarre tutti i dati per preparare la richiesta di booking alla compagnia di navigazione.

EMAIL:
{email_data.get("corpo", "")[:3000]}
{testo_allegati[:6000]}

Rispondi SOLO con JSON valido:
{{
  "tipo_richiesta": "ordine_spedizione",
  "cliente": {{"nome": "...", "paese": "IT", "email": null, "piva": null, "telefono": null}},
  "shipper": {{"nome": null, "paese": "CN", "indirizzo": null}},
  "consegnatario": {{"nome": null, "indirizzo": null, "citta": null}},
  "spedizione": {{
    "booking_number": null,
    "n_container": null,
    "tipo_container": null,
    "peso_totale_kg": null,
    "descrizione_merce": null,
    "hs_codes": [],
    "porto_carico": null,
    "porto_scarico": "NAPOLI",
    "nave_richiesta": null,
    "data_carico_richiesta": null,
    "eta_italia_richiesta": null
  }},
  "compagnia_preferita": null,
  "note_cliente": null,
  "urgenza": "alta",
  "richiesta_booking_carrier": "testo della email da inviare alla compagnia di navigazione per aprire il booking, in inglese, professionale"
}}""", max_tokens=2500)
    
    # Segna come processata
    mark_email_processed(body.uid)
    
    return {
        "step": 1,
        "stato": "ordine_analizzato",
        "dati_estratti": parse_json_safe(risposta),
        "operatore": user["sub"],
        "timestamp": datetime.now().isoformat()
    }


class SendEmailInput(BaseModel):
    to: str
    subject: str
    body: str
    html_body: Optional[str] = None

@router.post("/email/send")
async def send_email(body: SendEmailInput, user=Depends(get_current_user)):
    """Invia una email (bozza approvata) via SMTP."""
    success = await send_email_smtp(body.to, body.subject, body.body, body.html_body)
    if not success:
        raise HTTPException(status_code=500, detail="Errore nell'invio email")
    return {"ok": True, "timestamp": datetime.now().isoformat()}


# ─── STEP 1: Analisi email → estrai dati spedizione ─────────────────────────

class EmailInput(BaseModel):
    testo_email: str
    allegati: Optional[list] = []

@router.post("/step1/leggi-email")
async def step1_leggi_email(body: EmailInput, user=Depends(get_current_user)):
    ESCLUDI = ("(contenuto non estraibile)", "[Errore estrazione")
    testo_allegati = ""
    for a in (body.allegati or []):
        if not isinstance(a, dict):
            continue
        testo = a.get("testo", "")
        nome = a.get("nome", "")
        if any(testo.startswith(e) for e in ESCLUDI):
            continue
        testo_allegati += f"\n\n--- ALLEGATO: {nome} ---\n{testo[:2000]}"

    risposta = ai(f"""Sei un operatore di PF Ship Srl, agenzia spedizioni di Napoli.
Hai ricevuto questa email da un cliente italiano o dal tuo agente in Cina con un ordine di spedizione.
Devi estrarre tutti i dati per preparare la richiesta di booking alla compagnia di navigazione.

EMAIL:
{body.testo_email[:3000]}
{testo_allegati[:6000]}

Rispondi SOLO con JSON valido:
{{
  "tipo_richiesta": "ordine_spedizione",
  "cliente": {{"nome": "...", "paese": "IT", "email": null, "piva": null, "telefono": null}},
  "shipper": {{"nome": null, "paese": "CN", "indirizzo": null}},
  "consegnatario": {{"nome": null, "indirizzo": null, "citta": null}},
  "spedizione": {{
    "booking_number": null,
    "n_container": null,
    "tipo_container": null,
    "peso_totale_kg": null,
    "descrizione_merce": null,
    "hs_codes": [],
    "porto_carico": null,
    "porto_scarico": "NAPOLI",
    "nave_richiesta": null,
    "data_carico_richiesta": null,
    "eta_italia_richiesta": null
  }},
  "compagnia_preferita": null,
  "note_cliente": null,
  "urgenza": "alta",
  "richiesta_booking_carrier": "testo della email da inviare alla compagnia di navigazione per aprire il booking, in inglese, professionale"
}}""", max_tokens=2500)

    return {
        "step": 1,
        "stato": "ordine_analizzato",
        "dati_estratti": parse_json_safe(risposta),
        "operatore": user["sub"],
        "timestamp": datetime.now().isoformat()
    }


# ─── STEP 2: Apri pratica + genera booking confirmation ─────────────────────

class AperturaPraticaInput(BaseModel):
    dati_spedizione: dict
    booking_number_confermato: Optional[str] = ""
    compagnia_navigazione: Optional[str] = ""
    nave_confermata: Optional[str] = ""
    eta_confermata: Optional[str] = ""
    note_operatore: Optional[str] = ""
    allegati: Optional[list] = []

@router.post("/step2/apri-pratica")
async def step2_apri_pratica(
    body: AperturaPraticaInput,
    db: AsyncSession = Depends(get_sqlite),
    mysql_db: Session = Depends(get_mysql),
    user=Depends(get_current_user)
):
    sp = body.dati_spedizione.get("spedizione", {})
    cl = body.dati_spedizione.get("cliente", {})
    sh = body.dati_spedizione.get("shipper", {})
    conseg = body.dati_spedizione.get("consegnatario", {})

    if body.booking_number_confermato:
        sp["booking_number"] = body.booking_number_confermato
    if body.compagnia_navigazione:
        body.dati_spedizione["compagnia_navigazione"] = body.compagnia_navigazione
    if body.nave_confermata:
        sp["nave"] = body.nave_confermata
    if body.eta_confermata:
        sp["eta_italia"] = body.eta_confermata

    # Cerca il cliente/consegnatario in gedoanag per arricchire la pratica
    gedoanag_row = _lookup_gedoanag(mysql_db, nome=cl.get("nome"), piva=cl.get("piva"))
    gedoanag_id = gedoanag_row["id"] if gedoanag_row else None
    # Usa email/P.IVA da MySQL se non presenti nell'ordine
    cliente_email = cl.get("email") or (gedoanag_row.get("email") if gedoanag_row else None)
    consignee_piva = cl.get("piva") or (gedoanag_row.get("part_iva") if gedoanag_row else None)

    ESCLUDI = ("(contenuto non estraibile)", "[Errore")
    allegati_utili = [
        a for a in (body.allegati or [])
        if isinstance(a, dict) and not any(a.get("testo","").startswith(e) for e in ESCLUDI)
        and not a.get("nome","").lower().endswith((".jpg",".png",".gif",".txt"))
    ]

    result = await db.execute(text("""
        INSERT INTO pratiche (
            cliente, paese_origine, shipper, n_container, tipo_container,
            peso_totale_kg, descrizione_merce, porto_carico, porto_scarico,
            nave, viaggio, compagnia_navigazione, bl_number, booking_number,
            eta_italia, etd_cina, urgenza, note, operatore, step_corrente,
            allegati_json, consignee, consignee_piva, gedoanag_id
        ) VALUES (
            :cliente,:paese,:shipper,:n_cont,:tipo_cont,:peso,:merce,:porto_c,:porto_s,
            :nave,:viaggio,:compagnia,:bl,:booking,:eta,:etd,:urgenza,:note,:operatore,2,
            :allegati,:consignee,:consignee_piva,:gedoanag_id
        )
    """), {
        "cliente": cl.get("nome",""), "paese": cl.get("paese",""),
        "shipper": sh.get("nome","") if sh else "",
        "n_cont": sp.get("n_container",0), "tipo_cont": sp.get("tipo_container",""),
        "peso": sp.get("peso_totale_kg",0), "merce": sp.get("descrizione_merce",""),
        "porto_c": sp.get("porto_carico",""), "porto_s": sp.get("porto_scarico","NAPOLI"),
        "nave": sp.get("nave",""), "viaggio": sp.get("viaggio",""),
        "compagnia": body.dati_spedizione.get("compagnia_navigazione",""),
        "bl": body.dati_spedizione.get("bl_number",""),
        "booking": sp.get("booking_number",""),
        "eta": sp.get("eta_italia",""), "etd": sp.get("etd_cina",""),
        "urgenza": body.dati_spedizione.get("urgenza","alta"),
        "note": body.note_operatore, "operatore": user["sub"],
        "allegati": json.dumps(allegati_utili, ensure_ascii=False),
        "consignee": conseg.get("nome",""),
        "consignee_piva": consignee_piva or "",
        "gedoanag_id": gedoanag_id,
    })
    await db.commit()
    pratica_id = result.lastrowid

    booking_conf = ai(f"""Sei un operatore di PF Ship Srl Napoli.
Genera la booking confirmation da inviare all'agente per la pratica #{pratica_id}.

Dati spedizione:
- Cliente: {cl.get('nome','')}
- Nave: {sp.get('nave','')} Viaggio: {sp.get('viaggio','')}
- Da: {sp.get('porto_carico','')} A: {sp.get('porto_scarico','NAPOLI')}
- Container: {sp.get('n_container','')} x {sp.get('tipo_container','')}
- ETD: {sp.get('etd_cina','')} ETA: {sp.get('eta_italia','')}
- Booking N°: {sp.get('booking_number','')}

Scrivi email professionale in inglese con oggetto e corpo.
Firma: PF Ship Srl - Import Dept""", max_tokens=500)

    return {
        "step": 2, "stato": "pratica_aperta",
        "pratica_id": pratica_id,
        "gedoanag_collegato": gedoanag_id is not None,
        "booking_confirmation": booking_conf,
        "timestamp": datetime.now().isoformat()
    }


# ─── STEP 3: Monitora nave AIS ───────────────────────────────────────────────

class MonitoraggioInput(BaseModel):
    pratica_id: int
    mmsi: Optional[str] = None
    nave: Optional[str] = None
    eta_dichiarata: Optional[str] = None

@router.post("/step3/monitora-nave")
async def step3_monitora_nave(
    body: MonitoraggioInput,
    db: AsyncSession = Depends(get_sqlite),
    user=Depends(get_current_user)
):
    p = await get_pratica(body.pratica_id, db)

    # I parametri espliciti sovrascrivono quelli della pratica
    nave = body.nave or p.get("nave_attuale") or p.get("nave") or "N/D"
    eta = body.eta_dichiarata or p.get("eta_italia") or "N/D"
    mmsi = body.mmsi

    ais_data = None
    if mmsi:
        async with httpx.AsyncClient(timeout=20) as client:
            try:
                resp = await client.get(f"{AIS_BASE_URL}/ais/mt/{mmsi}/location/latest")
                ais_data = resp.json()
            except:
                ais_data = None

    if not ais_data:
        ais_data = {
            "mock": True, "mmsi": mmsi or "N/D",
            "vessel_name": nave,
            "lat": 36.5, "lon": 14.8,
            "speed": 18.4, "destination": "NAPLES",
            "eta": eta,
            "status": "Under way using engine",
            "distance_to_naples_nm": 420
        }

    analisi = ai(f"""Operatore doganale PF Ship Srl. Pratica #{body.pratica_id}.
Nave: {nave} - ETA dichiarata: {eta}
Dati AIS: {json.dumps(ais_data)}

Calcola e rispondi in italiano:
1. Giorni mancanti all'arrivo a Napoli
2. Entro quando richiedere i documenti al cliente (2 sett. prima ETA)
3. Entro quando fare richiesta delivery order a Maersk
4. Segnala se c'è ritardo rispetto all'ETA dichiarata
Sii conciso.""", max_tokens=300)

    await _crea_approvazione(db, body.pratica_id, 3, "VesselTracker", analisi)

    return {
        "step": 3, "stato": "nave_monitorata",
        "pratica_id": body.pratica_id,
        "nave": nave, "eta": eta,
        "ais": ais_data, "analisi": analisi,
        "timestamp": datetime.now().isoformat()
    }


# ─── STEP 4: Richiesta documenti al cliente ──────────────────────────────────

class RichiestaDocInput(BaseModel):
    pratica_id: int
    lingua: Optional[str] = "italiano"
    note_extra: Optional[str] = ""

@router.post("/step4/richiedi-documenti")
async def step4_richiedi_documenti(
    body: RichiestaDocInput,
    db: AsyncSession = Depends(get_sqlite),
    mysql_db: Session = Depends(get_mysql),
    user=Depends(get_current_user)
):
    p = await get_pratica(body.pratica_id, db)

    cliente_nome = p.get("cliente","")
    eta_nave = p.get("eta_italia","")
    n_container = p.get("n_container",0)

    # Email cliente: cerca in gedoanag se disponibile
    cliente_email = ""
    if p.get("gedoanag_id"):
        anag = _lookup_gedoanag(mysql_db, piva=p.get("consignee_piva"))
        if not anag:
            anag = _lookup_gedoanag(mysql_db, nome=cliente_nome)
        cliente_email = (anag or {}).get("email","") or ""

    email_richiesta = ai(f"""PF Ship Srl Napoli, pratica #{body.pratica_id}.
Genera email professionale in {body.lingua} per richiedere documentazione di sdoganamento a:
Cliente: {cliente_nome}{(' - Email: ' + cliente_email) if cliente_email else ''}
ETA nave a Napoli: {eta_nave}
N° container: {n_container}
{('Note: ' + body.note_extra) if body.note_extra else ''}

Documenti da richiedere:
1. Fattura commerciale (in inglese, con valori doganali)
2. Packing list dettagliata
3. Polizza di carico (B/L originale o copia)
4. Packing declaration
5. Certificato di origine (se richiesto)

Spiega l'urgenza (arrivo nave tra circa 2 settimane).
Oggetto + corpo email. Firma: Ufficio Import PF Ship Srl""", max_tokens=500)

    await _crea_approvazione(db, body.pratica_id, 4, "DocRequest", email_richiesta)

    return {
        "step": 4, "stato": "documenti_richiesti",
        "pratica_id": body.pratica_id,
        "cliente": cliente_nome,
        "email_cliente": cliente_email,
        "bozza_email": email_richiesta,
        "timestamp": datetime.now().isoformat()
    }


# ─── STEP 5: Controllo documenti ricevuti ────────────────────────────────────

class DocumentiInput(BaseModel):
    pratica_id: int
    documenti_ricevuti: list[str]
    note_anomalie: Optional[str] = ""
    lingua: Optional[str] = "italiano"

@router.post("/step5/controlla-documenti")
async def step5_controlla_documenti(
    body: DocumentiInput,
    db: AsyncSession = Depends(get_sqlite),
    mysql_db: Session = Depends(get_mysql),
    user=Depends(get_current_user)
):
    p = await get_pratica(body.pratica_id, db)

    cliente_nome = p.get("cliente","")
    eta_nave = p.get("eta_italia","")

    cliente_email = ""
    if p.get("gedoanag_id"):
        anag = _lookup_gedoanag(mysql_db, nome=cliente_nome)
        cliente_email = (anag or {}).get("email","") or ""

    obbligatori = ["Fattura commerciale", "Packing list", "Polizza di carico (B/L)"]
    raccomandati = ["Packing declaration", "Certificato di origine"]

    mancanti_obb = [d for d in obbligatori if d not in body.documenti_ricevuti]
    mancanti_racc = [d for d in raccomandati if d not in body.documenti_ricevuti]
    stato = "BLOCCATO" if mancanti_obb else ("INCOMPLETO" if mancanti_racc else "OK")

    destinatario = f"{cliente_nome}{(' (' + cliente_email + ')') if cliente_email else ''}"
    azione = ai(f"""Esperto doganale PF Ship Srl. Pratica #{body.pratica_id} - {cliente_nome}
ETA nave: {eta_nave}
Documenti ricevuti: {', '.join(body.documenti_ricevuti) or 'Nessuno'}
Mancanti obbligatori: {', '.join(mancanti_obb) or 'Nessuno'}
Anomalie: {body.note_anomalie or 'Nessuna'}
Stato: {stato}

{"Genera email di sollecito/correzione in " + body.lingua + " per " + destinatario + " con oggetto e corpo." if stato != "OK" else "Conferma completezza e indica prossimi step operativi (bolla doganale)."}""", max_tokens=600)

    await _crea_approvazione(db, body.pratica_id, 5, "DocCheck", azione)

    return {
        "step": 5, "stato": stato,
        "pratica_id": body.pratica_id,
        "cliente": cliente_nome,
        "email_cliente": cliente_email,
        "mancanti_obbligatori": mancanti_obb,
        "mancanti_raccomandati": mancanti_racc,
        "azione": azione,
        "timestamp": datetime.now().isoformat()
    }


# ─── STEP 6: Bolla doganale ──────────────────────────────────────────────────

class BollaInput(BaseModel):
    pratica_id: int

@router.post("/step6/bolla-doganale")
async def step6_bolla_doganale(body: BollaInput, db: AsyncSession = Depends(get_sqlite), user=Depends(get_current_user)):
    p = await get_pratica(body.pratica_id, db)

    allegati = []
    if p.get("allegati_json"):
        try:
            allegati = json.loads(p["allegati_json"])
        except:
            allegati = []

    testo_allegati = "\n\n".join(
        f"--- {a['nome']} ---\n{a['testo'][:3000]}"
        for a in allegati if a.get("testo")
    )

    # Usa nave_attuale se presente (trasbordo), altrimenti nave originale
    nave_trasporto = p.get("nave_attuale") or p.get("nave","")

    risposta = ai(f"""Sei un esperto doganale italiano. Devi pre-compilare la bolla doganale per PRADO.

DATI PRATICA:
- Cliente/Importatore: {p.get('consignee', p.get('cliente',''))} — P.IVA: {p.get('consignee_piva','')}
- Shipper/Esportatore: {p.get('shipper','')}
- N° container: {p.get('n_container','')} x {p.get('tipo_container','')} — {p.get('container_number','')} sigillo {p.get('seal_number','')}
- Nave: {nave_trasporto} - Viaggio: {p.get('viaggio','')}
- Porto carico: {p.get('porto_carico','')} → Porto scarico: {p.get('porto_scarico','')}
- MBL: {p.get('bl_number','')} — HBL: {p.get('hbl_number','')}
- Invoice: {p.get('invoice_number','')} — Valore FOB: EUR {p.get('valore_merce_eur',0)}
- Peso totale lordo: {p.get('peso_totale_kg','')} kg — Peso netto: {p.get('peso_netto_kg','')} kg
- N° colli: {p.get('n_colli','')} — N° pezzi: {p.get('n_pezzi','')}
- Data imbarco: {p.get('etd_cina','')} — ETA Napoli: {p.get('eta_italia','')}

DOCUMENTI ALLEGATI:
{testo_allegati[:8000]}

Estrai dai documenti e restituisci SOLO un JSON con questa struttura:
{{
  "importatore": {{
    "ragione_sociale": "...",
    "piva": "...",
    "indirizzo": "...",
    "citta": "..."
  }},
  "esportatore": {{
    "ragione_sociale": "...",
    "paese": "CN"
  }},
  "trasporto": {{
    "mezzo": "nave",
    "nome_nave": "...",
    "viaggio": "...",
    "bl_number": "...",
    "porto_carico": "...",
    "porto_scarico": "...",
    "data_imbarco": "..."
  }},
  "merci": [
    {{
      "descrizione": "...",
      "codice_hs": "...",
      "quantita": "...",
      "unita": "CTN",
      "peso_lordo_kg": 0,
      "peso_netto_kg": 0,
      "valore_eur": 0,
      "dazio_pct": 0,
      "paese_origine": "CN"
    }}
  ],
  "totali": {{
    "colli_totali": 0,
    "peso_lordo_kg": 0,
    "peso_netto_kg": 0,
    "valore_totale_eur": 0,
    "imponibile_dazi": 0,
    "totale_dazi": 0,
    "iva_pct": 22,
    "totale_iva": 0
  }},
  "note_prado": "istruzioni specifiche per compilare PRADO"
}}""", max_tokens=3000)

    bolla = parse_json_safe(risposta)
    await _crea_approvazione(db, body.pratica_id, 6, "CustomsFiler", json.dumps(bolla, default=str))

    return {
        "step": 6,
        "stato": "bolla_precompilata",
        "pratica_id": body.pratica_id,
        "bolla": bolla,
        "timestamp": datetime.now().isoformat()
    }


# ─── STEP 7: Richiesta delivery order ────────────────────────────────────────

class DeliveryOrderInput(BaseModel):
    pratica_id: int
    note: Optional[str] = ""

@router.post("/step7/delivery-order")
async def step7_delivery_order(
    body: DeliveryOrderInput,
    db: AsyncSession = Depends(get_sqlite),
    user=Depends(get_current_user)
):
    p = await get_pratica(body.pratica_id, db)

    bl_number = p.get("bl_number","")
    booking_number = p.get("booking_number","")
    compagnia = p.get("compagnia_navigazione","")
    cliente = p.get("cliente","")
    n_container = p.get("n_container",0)

    istruzioni = ai(f"""Operatore PF Ship Srl. Pratica #{body.pratica_id}.
Richiesta delivery order a {compagnia}:
B/L: {bl_number} | Booking: {booking_number}
Cliente: {cliente} | Container: {n_container}
{('Note: ' + body.note) if body.note else ''}

Genera:
1. Testo email/richiesta formale a {compagnia} per il delivery order
2. Checklist pre-richiesta (fatture da pagare, documenti necessari)
3. Cosa verificare prima di inviare la richiesta
4. Tempi attesi di risposta da {compagnia}
Professionale e specifico.""", max_tokens=600)

    await _crea_approvazione(db, body.pratica_id, 7, "DOAgent", istruzioni)

    return {
        "step": 7, "stato": "delivery_order_richiesto",
        "pratica_id": body.pratica_id,
        "bl_number": bl_number,
        "compagnia": compagnia,
        "istruzioni": istruzioni,
        "timestamp": datetime.now().isoformat()
    }


# ─── STEP 8: Ricezione e pagamento fatture compagnia ─────────────────────────

class PagamentoInput(BaseModel):
    pratica_id: int
    importo_fattura: Optional[float] = 0
    valuta: Optional[str] = "USD"
    note: Optional[str] = ""

@router.post("/step8/pagamento-fatture")
async def step8_pagamento(
    body: PagamentoInput,
    db: AsyncSession = Depends(get_sqlite),
    user=Depends(get_current_user)
):
    p = await get_pratica(body.pratica_id, db)
    compagnia = p.get("compagnia_navigazione","")

    istruzioni = ai(f"""Operatore contabile PF Ship Srl. Pratica #{body.pratica_id}.
Ricezione fattura da {compagnia}: {body.valuta} {body.importo_fattura}
Note: {body.note or 'Nessuna'}

Genera checklist per:
1. Verifica fattura (voci, importi, riferimenti B/L)
2. Procedura di pagamento internazionale
3. Registrazione in contabilità
4. Dopo il pagamento: cosa richiedere a {compagnia}
Conciso e operativo.""", max_tokens=400)

    await _crea_approvazione(db, body.pratica_id, 8, "PaymentAgent", istruzioni)

    return {
        "step": 8, "stato": "pagamento_gestito",
        "pratica_id": body.pratica_id,
        "compagnia": compagnia,
        "istruzioni": istruzioni,
        "timestamp": datetime.now().isoformat()
    }


# ─── STEP 9: Prenotazione trasportatore ──────────────────────────────────────

class TrasportatoreInput(BaseModel):
    pratica_id: int
    note: Optional[str] = ""

@router.post("/step9/prenota-trasportatore")
async def step9_trasportatore(
    body: TrasportatoreInput,
    db: AsyncSession = Depends(get_sqlite),
    mysql_db: Session = Depends(get_mysql),
    user=Depends(get_current_user)
):
    p = await get_pratica(body.pratica_id, db)

    cliente_nome = p.get("cliente","")
    n_container = p.get("n_container",0)
    eta_nave = p.get("eta_italia","")

    # Indirizzo consegna: da gedoanag se disponibile, fallback da pratica
    cliente_email = ""
    indirizzo_consegna = ""
    if p.get("gedoanag_id"):
        anag = _lookup_gedoanag(mysql_db, piva=p.get("consignee_piva"), nome=cliente_nome)
        if anag:
            cliente_email = anag.get("email","") or ""
            loc = anag.get("localita","")
            ind = anag.get("indirizzo","")
            indirizzo_consegna = f"{ind}, {loc}".strip(", ")
    if not indirizzo_consegna:
        # Fallback: consignee dalla pratica
        consignee_nome = p.get("consignee","")
        if consignee_nome:
            anag2 = _lookup_gedoanag(mysql_db, nome=consignee_nome)
            if anag2:
                cliente_email = cliente_email or anag2.get("email","") or ""
                indirizzo_consegna = f"{anag2.get('indirizzo','')}, {anag2.get('localita','')}".strip(", ")

    email_bozze = ai(f"""PF Ship Srl. Pratica #{body.pratica_id}.
Ricezione delivery order confermata. Contatta il cliente per confermare data consegna.

Cliente: {cliente_nome}{(' (' + cliente_email + ')') if cliente_email else ''}
Indirizzo consegna: {indirizzo_consegna or 'da confermare con cliente'}
N° container: {n_container}
ETA nave: {eta_nave}
{('Note: ' + body.note) if body.note else ''}

Genera:
1. Email al cliente per concordare data/ora consegna
2. Informazioni che il cliente deve confermare
3. Email al trasportatore per prenotare il ritiro dal porto
Tutto in italiano, professionale.""", max_tokens=600)

    await _crea_approvazione(db, body.pratica_id, 9, "TransportAgent", email_bozze)

    return {
        "step": 9, "stato": "trasportatore_prenotato",
        "pratica_id": body.pratica_id,
        "cliente": cliente_nome,
        "email_cliente": cliente_email,
        "indirizzo_consegna": indirizzo_consegna,
        "bozze_email": email_bozze,
        "timestamp": datetime.now().isoformat()
    }


# ─── STEP 10: Conferma consegna e chiusura operativa ─────────────────────────

class ConsegnaInput(BaseModel):
    pratica_id: int
    data_consegna: Optional[str] = ""
    note_consegna: Optional[str] = ""

@router.post("/step10/conferma-consegna")
async def step10_consegna(
    body: ConsegnaInput,
    db: AsyncSession = Depends(get_sqlite),
    user=Depends(get_current_user)
):
    p = await get_pratica(body.pratica_id, db)
    cliente = p.get("cliente","")

    riepilogo = ai(f"""PF Ship Srl. Pratica #{body.pratica_id} - consegna effettuata.
Cliente: {cliente}
Data consegna: {body.data_consegna}
Note: {body.note_consegna or 'Nessuna'}

Genera:
1. Email di conferma consegna al cliente
2. Checklist chiusura operativa (documenti da archiviare)
3. Istruzioni per passare pratica alla contabilità per fatturazione
In italiano, professionale.""", max_tokens=500)

    await _aggiorna_step(db, body.pratica_id, 10, extra={"stato": "consegnata"})
    await _crea_approvazione(db, body.pratica_id, 10, "DeliveryAgent", riepilogo)

    return {
        "step": 10, "stato": "consegna_confermata",
        "pratica_id": body.pratica_id,
        "cliente": cliente,
        "riepilogo": riepilogo,
        "timestamp": datetime.now().isoformat()
    }


# ─── STEP 11: Fatturazione (Ge.FA) ───────────────────────────────────────────

class FatturazioneInput(BaseModel):
    pratica_id: int
    data_consegna: Optional[str] = ""

@router.post("/step11/fatturazione")
async def step11_fatturazione(
    body: FatturazioneInput,
    db: AsyncSession = Depends(get_sqlite),
    mysql_db: Session = Depends(get_mysql),
    user=Depends(get_current_user)
):
    p = await get_pratica(body.pratica_id, db)

    cliente = p.get("cliente","")
    n_container = p.get("n_container",0)
    tipo_container = p.get("tipo_container","") or "40HC"
    descrizione_merce = p.get("descrizione_merce","")
    bl_number = p.get("bl_number","")

    tariffe = []
    try:
        rows = mysql_db.execute(text("""
            SELECT desc_clien, tipo_fattu, impo_impon, importo, anno_fattu
            FROM gefadcts
            WHERE desc_clien LIKE :cliente
              AND (cancellato = 0 OR cancellato IS NULL)
            ORDER BY anno_fattu DESC, id DESC
            LIMIT 5
        """), {"cliente": f"%{cliente[:20]}%"}).fetchall()
        tariffe = [dict(r._mapping) for r in rows]
    except:
        pass

    istruzioni = ai(f"""Contabile PF Ship Srl. Pratica #{body.pratica_id}.
Cliente: {cliente}
Spedizione: {n_container} container {tipo_container} - {descrizione_merce}
B/L: {bl_number}
Data consegna: {body.data_consegna}

Fatture precedenti trovate in Ge.FA: {json.dumps(tariffe, default=str) if tariffe else 'Nessuna trovata'}

Genera:
1. Voci da fatturare tipiche per una spedizione import da Cina ({n_container} container {tipo_container})
   (nolo marittimo, spese portuali, dogana, handling, trasporto, diritti)
2. Istruzioni per emettere fattura su Ge.FA
3. Riferimenti da indicare in fattura (B/L, n° pratica, ecc.)
In italiano.""", max_tokens=600)

    await _crea_approvazione(db, body.pratica_id, 11, "InvoiceAgent", istruzioni)

    return {
        "step": 11, "stato": "fatturazione_preparata",
        "pratica_id": body.pratica_id,
        "cliente": cliente,
        "fatture_precedenti": tariffe,
        "istruzioni": istruzioni,
        "timestamp": datetime.now().isoformat()
    }


# ─── STEP 12: Registrazione contabilità (Ge.CO) ──────────────────────────────

class ContabilitaInput(BaseModel):
    pratica_id: int
    importo_fattura: Optional[float] = 0
    numero_fattura: Optional[str] = ""

@router.post("/step12/registra-contabilita")
async def step12_contabilita(
    body: ContabilitaInput,
    db: AsyncSession = Depends(get_sqlite),
    mysql_db: Session = Depends(get_mysql),
    user=Depends(get_current_user)
):
    p = await get_pratica(body.pratica_id, db)

    cliente = p.get("cliente","")
    invoice_number = p.get("invoice_number","") or body.numero_fattura

    # Cerca in Ge.CO: prima per P.IVA, poi per nome
    soggetto = None
    try:
        consignee_piva = p.get("consignee_piva","")
        if consignee_piva:
            row = mysql_db.execute(text("""
                SELECT ragi_socia, codi_conto, indirizzo, localita
                FROM cogeanag
                WHERE codi_fisca = :piva OR codi_fisca LIKE :piva2
                  AND (cancellato = 0 OR cancellato IS NULL)
                LIMIT 1
            """), {"piva": consignee_piva, "piva2": f"%{consignee_piva}%"}).fetchone()
            if row:
                soggetto = dict(row._mapping)
        if not soggetto:
            row = mysql_db.execute(text("""
                SELECT ragi_socia, codi_conto, indirizzo, localita
                FROM cogeanag
                WHERE ragi_socia LIKE :nome
                  AND (cancellato = 0 OR cancellato IS NULL)
                LIMIT 1
            """), {"nome": f"%{cliente[:20]}%"}).fetchone()
            if row:
                soggetto = dict(row._mapping)
    except:
        pass

    istruzioni = ai(f"""Contabile PF Ship Srl. Chiusura pratica #{body.pratica_id}.
Cliente: {cliente}
Importo fattura: €{body.importo_fattura} | N° fattura: {invoice_number}
Dati Ge.CO: {json.dumps(soggetto, default=str) if soggetto else 'Cliente da verificare in anagrafica'}

Genera istruzioni per:
1. Registrazione fattura attiva in Ge.CO
2. Piano dei conti da usare
3. Scadenze di pagamento da monitorare
4. Chiusura definitiva pratica e archiviazione
In italiano, operativo.""", max_tokens=500)

    await _aggiorna_step(db, body.pratica_id, 12, extra={"stato": "chiusa"})
    await _crea_approvazione(db, body.pratica_id, 12, "AccountingAgent", istruzioni)

    return {
        "step": 12, "stato": "pratica_chiusa",
        "pratica_id": body.pratica_id,
        "cliente": cliente,
        "soggetto_coge": soggetto,
        "istruzioni": istruzioni,
        "timestamp": datetime.now().isoformat()
    }


# ─── Aggiorna step corrente ──────────────────────────────────────────────────

class AggiornaSteInput(BaseModel):
    pratica_id: int
    step: int

@router.post("/aggiorna-step")
async def aggiorna_step(body: AggiornaSteInput, db: AsyncSession = Depends(get_sqlite), user=Depends(get_current_user)):
    await db.execute(text(
        "UPDATE pratiche SET step_corrente = :step WHERE id = :id"
    ), {"step": body.step, "id": body.pratica_id})
    await db.commit()
    return {"ok": True}


# ─── Lista pratiche ───────────────────────────────────────────────────────────

@router.get("/pratiche")
async def lista_pratiche(db: AsyncSession = Depends(get_sqlite), user=Depends(get_current_user)):
    try:
        result = await db.execute(text("SELECT * FROM pratiche ORDER BY creata_il DESC"))
        rows = result.fetchall()
        return [dict(r._mapping) for r in rows]
    except:
        return []
